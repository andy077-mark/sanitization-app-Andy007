"""Background job processing, packaging, audit report generation, and failure handling."""
from __future__ import annotations

import json
import logging
import os
import shutil
import tempfile
import time
import zipfile
from datetime import datetime
from pathlib import Path

from . import config as cfg
from . import rules as rules_module
from .sanitize import (
    archive_stem,
    audit_add,
    classify_rule,
    extract_archive,
    get_archive_ext,
    is_archive_name,
    repack_directory,
    sanitize_directory,
    sanitize_file_if_supported,
    sanitize_relative_path,
)


def _keyword_map(rules: list[str] | None) -> dict[str, list[str]]:
    """Map audit match types to the configured rules/keywords that produced them.

    The report intentionally shows the configured rule text exactly as entered.
    It never attempts to reconstruct the original matched value from sanitized data.
    """
    mapping: dict[str, list[str]] = {}
    for idx, rule in enumerate(rules or [], 1):
        kind = classify_rule(rule, idx)
        mapping.setdefault(kind, []).append(rule)
    return mapping


def _keywords_for_match(match_type: str, mapping: dict[str, list[str]]) -> str:
    """Return configured keyword/rule text for one audit match type."""
    base_type = str(match_type or "")
    if ": " in base_type:
        base_type = base_type.split(": ", 1)[1]

    if base_type == "Non-ASCII Character":
        return "Non-ASCII Character"
    if base_type == "Filename Format":
        return "Invalid filename character"

    keywords = mapping.get(base_type, [])
    if keywords:
        return " | ".join(keywords)
    return base_type


def create_audit_report(
    job_id: str,
    stats: list[dict],
    audit: dict,
    total_replacements: int,
    duration: float,
    status: str = "Done",
    error: str | None = None,
    rules: list[str] | None = None,
) -> Path:
    """Create the XLSX audit report with configured keywords shown unsanitized.

    Sanitized output files remain sanitized. The Keywords column contains only
    the configured rule/keyword text, not a reconstructed original matched value.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    report = cfg.OUTPUT / f"Sanitization_Report_{job_id}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["SOC Data Sanitization Report"])
    ws.append(["Job ID", job_id])
    ws.append(["Status", status])
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws.append(["Total Replacements", total_replacements])
    ws.append(["Duration (seconds)", round(duration, 3)])
    if error:
        ws.append(["Error", error])
    ws.append([])
    ws.append(["File Name", "Replacements", "Duration (seconds)"])
    for item in stats:
        ws.append([item.get("file", ""), item.get("replacements", 0), item.get("duration", 0)])

    keyword_map = _keyword_map(rules)
    audit_ws = wb.create_sheet("Audit")
    audit_ws.append(["File Name", "Match Type", "Occurrences", "Location", "Keywords"])
    for item in sorted(audit.values(), key=lambda x: (x["file"], x["match_type"])):
        audit_ws.append(
            [
                item["file"],
                item["match_type"],
                item["occurrences"],
                item["location"],
                _keywords_for_match(item["match_type"], keyword_map),
            ]
        )
    if not audit:
        audit_ws.append(["", "No sensitive matches detected", 0, "", ""])

    for sheet in (ws, audit_ws):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="123447")
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 70)
            sheet.column_dimensions[column[0].column_letter].width = max(12, width)
    wb.save(report)
    return report


def _session_set(job_id: str, **fields) -> None:
    info = cfg.SESSIONS.setdefault(job_id, {})
    info.update(fields)


def _zip_compression_level() -> int:
    """Return a speed-oriented ZIP level while allowing local tuning.

    Level 1 is deliberately the default: SOC log bundles favor fast completion
    over maximum compression. Set SANIT_ZIP_COMPRESSION_LEVEL=0..9 if required.
    """
    try:
        value = int(os.environ.get("SANIT_ZIP_COMPRESSION_LEVEL", "1"))
    except ValueError:
        value = 1
    return max(0, min(value, 9))


def _package_outputs(job_id: str, sanitized_root: Path, report: Path) -> Path:
    """Create the downloadable package directly from sanitized outputs.

    The previous implementation copied the entire sanitized tree into another
    temporary Bundle directory before ZIP creation. For large SOC logs that
    doubled disk I/O. Writing the ZIP directly avoids that extra full-file copy.
    """
    bundle = cfg.OUTPUT / f"Sanitized_Package_{job_id}.zip"
    level = _zip_compression_level()
    with zipfile.ZipFile(
        bundle,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=level,
        allowZip64=True,
    ) as archive:
        for path in sanitized_root.rglob("*"):
            if not path.is_file():
                continue
            relative = path.relative_to(sanitized_root).as_posix()
            archive.write(path, arcname=f"Sanitized_Files/{relative}")
        archive.write(report, arcname=report.name)
    return bundle


def process_file(
    job_id: str,
    upload_paths: list[Path],
    original_names: list[str],
    relative_paths: list[str],
    ip: str | None,
    rules_override: list[str] | None = None,
) -> None:
    """Process all uploaded files as one job and produce one ZIP package + Excel report."""
    cfg.sem.acquire()
    started = datetime.now()
    perf_started = time.perf_counter()
    job_tmp: Path | None = None
    try:
        _session_set(job_id, status="Processing", progress=5)
        cfg.db_update_job(job_id, status="Processing")
        active_rules = rules_override if rules_override is not None else rules_module.reload_bad_patterns()
        job_tmp = Path(tempfile.mkdtemp(prefix=f"job_{job_id}_", dir=cfg.TEMP))
        sanitized_root = job_tmp / "Sanitized_Files"
        sanitized_root.mkdir()

        audit: dict = {}
        stats: list[dict] = []
        total_replacements = 0
        file_count = max(1, len(upload_paths))

        for idx, up in enumerate(upload_paths):
            original = original_names[idx]
            rel_raw = relative_paths[idx] if idx < len(relative_paths) else original
            rel_safe, path_repl, path_counts = sanitize_relative_path(rel_raw, active_rules)
            if len(rel_safe.parts) > 50:
                raise ValueError("Folder path is too deep")
            target_parent = sanitized_root / rel_safe.parent
            target_parent.mkdir(parents=True, exist_ok=True)
            safe_name = rel_safe.name
            audit_name = rel_safe.as_posix()
            replacements = path_repl
            if path_repl:
                for kind, count in path_counts.items():
                    audit_add(audit, audit_name, f"Path: {kind}", count, "Path", rel_safe.as_posix())

            file_started = time.perf_counter()
            if is_archive_name(original):
                original_ext = get_archive_ext(original)
                extracted = Path(tempfile.mkdtemp(prefix="archive_", dir=job_tmp))
                try:
                    extract_archive(up, extracted, 0)
                    replacements += sanitize_directory(extracted, active_rules, audit, audit_name, 0)
                    out_base = target_parent / archive_stem(safe_name)
                    out = repack_directory(extracted, out_base, original_ext)
                finally:
                    shutil.rmtree(extracted, ignore_errors=True)
            else:
                candidate = target_parent / safe_name
                counter = 1
                while candidate.exists():
                    stem, ext = os.path.splitext(safe_name)
                    candidate = target_parent / f"{stem}_{counter}{ext}"
                    counter += 1

                # Uploads and the job workspace live under the same application
                # filesystem. Moving is normally a metadata operation and avoids
                # a complete extra copy of large logs before sanitization.
                try:
                    shutil.move(str(up), str(candidate))
                except OSError:
                    # Safe fallback for unusual cross-filesystem deployments.
                    shutil.copy2(up, candidate)
                replacements += sanitize_file_if_supported(candidate, active_rules, audit, audit_name)
                out = candidate

            elapsed = time.perf_counter() - file_started
            try:
                output_rel = out.relative_to(sanitized_root).as_posix()
            except ValueError:
                output_rel = rel_safe.as_posix()
            stats.append({"file": output_rel, "replacements": replacements, "duration": round(elapsed, 5)})
            total_replacements += replacements
            progress = 10 + int(((idx + 1) / file_count) * 70)
            _session_set(job_id, progress=progress, current_file=rel_raw)

        processing_elapsed = time.perf_counter() - perf_started
        _session_set(job_id, progress=85, current_file="Generating audit report")
        report_started = time.perf_counter()
        duration_so_far = (datetime.now() - started).total_seconds()
        report = create_audit_report(
            job_id,
            stats,
            audit,
            total_replacements,
            duration_so_far,
            rules=active_rules,
        )
        report_elapsed = time.perf_counter() - report_started

        _session_set(job_id, progress=92, current_file="Packaging sanitized outputs")
        package_started = time.perf_counter()
        bundle = _package_outputs(job_id, sanitized_root, report)
        package_elapsed = time.perf_counter() - package_started

        duration = (datetime.now() - started).total_seconds()
        token = cfg.SESSIONS[job_id]["token"]
        completed = datetime.now().isoformat(timespec="seconds")
        session_update = {
            "status": "Done",
            "progress": 100,
            "file_stats": stats,
            "total_replacements": total_replacements,
            "duration": duration,
            "output_path": str(bundle),
            "report_path": str(report),
            "download": f"/download/{job_id}?token={token}",
            "report_download": f"/download/report/{job_id}?token={token}",
            "time": datetime.now(),
        }
        _session_set(job_id, **session_update)
        cfg.db_update_job(
            job_id,
            status="Done",
            completed_at=completed,
            total_replacements=total_replacements,
            duration=duration,
            output_path=str(bundle),
            report_path=str(report),
            error=None,
        )

        with open(cfg.LOGS / f"{job_id}.log", "w", encoding="utf-8") as handle:
            handle.write(
                json.dumps(
                    {
                        "job_id": job_id,
                        "status": "Done",
                        "files": [item["file"] for item in stats],
                        "time": completed,
                        "ip": ip,
                        "duration": round(duration, 3),
                        "replacements": total_replacements,
                        "performance": {
                            "processing_seconds": round(processing_elapsed, 3),
                            "report_seconds": round(report_elapsed, 3),
                            "packaging_seconds": round(package_elapsed, 3),
                            "zip_compression_level": _zip_compression_level(),
                        },
                    }
                )
                + "\n"
            )

    except Exception as exc:
        duration = (datetime.now() - started).total_seconds()
        message = str(exc) or exc.__class__.__name__
        logging.exception("Job %s failed", job_id)
        _session_set(
            job_id,
            status="Failed",
            progress=100,
            error=message,
            duration=duration,
            time=datetime.now(),
        )
        cfg.db_update_job(
            job_id,
            status="Failed",
            completed_at=datetime.now().isoformat(timespec="seconds"),
            duration=duration,
            error=message,
        )
        try:
            with open(cfg.LOGS / f"{job_id}.log", "w", encoding="utf-8") as handle:
                handle.write(json.dumps({"job_id": job_id, "status": "Failed", "error": message}) + "\n")
        except OSError:
            pass
    finally:
        if job_tmp:
            shutil.rmtree(job_tmp, ignore_errors=True)
        for up in upload_paths:
            try:
                Path(up).unlink(missing_ok=True)
            except OSError:
                pass
        cfg.sem.release()
