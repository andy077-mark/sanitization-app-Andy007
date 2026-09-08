"""Core sanitization, audit capture, archive safety, and repacking helpers."""
from __future__ import annotations

import bz2
import gzip
import io
import lzma
import os
import re
import shutil
import stat
import subprocess
import tarfile
import tempfile
import uuid
import zipfile
from pathlib import Path, PurePosixPath

from . import config as cfg
from . import rules as rules_module
from .document_sanitizers import sanitize_docx_file as sanitize_docx_document
from .document_sanitizers import sanitize_pdf_file as sanitize_pdf_document
from .text_encoding import TEXT_LIKE_EXTS, detect_text_encoding, is_known_binary_extension


def x_replacer(match):
    return "X" * len(match.group(0))


def classify_rule(pattern: str, index: int) -> str:
    p = pattern.lower()
    if "@" in p or "email" in p:
        return "Email Address"
    if "ipv4" in p or (r"\d{1,3}" in p and r"\." in p):
        return "IPv4 Address"
    if "ipv6" in p or ("[0-9a-f" in p and ":" in p):
        return "IPv6 Address"
    if "mac" in p or (":" in p and "{6}" in p):
        return "MAC Address"
    if (
        "password" in p or "passwd" in p or "secret" in p or "token" in p
        or ("api" in p and "key" in p)
    ):
        return "Credential / Secret"
    if "http" in p or "url" in p:
        return "URL"
    if "host" in p or "domain" in p:
        return "Hostname / Domain"
    if "user" in p or "employee" in p:
        return "Identity"
    return f"Custom Rule {index}"


def audit_add(audit: dict, file_name: str, match_type: str, count: int, location: str, safe_example: str) -> None:
    """Aggregate an audit row without ever storing the original sensitive value."""
    if count <= 0:
        return
    key = (file_name, match_type)
    item = audit.setdefault(
        key,
        {
            "file": file_name,
            "match_type": match_type,
            "occurrences": 0,
            "location": location,
            "safe_example": "",
        },
    )
    item["occurrences"] += count
    if not item["safe_example"] and safe_example:
        example = " ".join(str(safe_example).strip().split())
        item["safe_example"] = example[: cfg.AUDIT_EXAMPLE_MAX]
        item["location"] = location


def sanitize_content(text: str, rules: list[str] | None = None):
    counts: dict[str, int] = {}
    total = 0
    text, n = re.subn(r"[^\x20-\x7E\n\r]", x_replacer, text)
    if n:
        counts["Non-ASCII Character"] = n
        total += n
    pats = rules if rules is not None else rules_module.BAD_PATTERNS
    for idx, pat in enumerate(pats, 1):
        text, n = re.subn(pat, x_replacer, text, flags=re.IGNORECASE)
        if n:
            kind = classify_rule(pat, idx)
            counts[kind] = counts.get(kind, 0) + n
            total += n
    return text, total, counts


def is_text(fp: Path) -> bool:
    """Return True only when the file can be decoded safely as supported text."""
    if is_known_binary_extension(fp):
        return False
    return detect_text_encoding(fp) is not None


def sanitize_text_file(
    fpath: Path,
    rules: list[str],
    audit: dict,
    audit_name: str,
    encoding_info: tuple[str, bytes] | None = None,
) -> int:
    """Sanitize a text file while preserving its detected character encoding/BOM."""
    detected = encoding_info or detect_text_encoding(fpath)
    if detected is None:
        raise ValueError(f"Unsupported text encoding or binary content: {fpath.name}")

    codec, bom = detected
    total = 0
    tmp_out = fpath.with_suffix(fpath.suffix + ".sanit.tmp")
    try:
        with open(fpath, "rb") as src_raw, open(tmp_out, "wb") as dst_raw:
            if bom:
                actual_bom = src_raw.read(len(bom))
                if actual_bom != bom:
                    raise ValueError(f"Text encoding marker changed while reading: {fpath.name}")
                dst_raw.write(bom)

            # newline="" keeps the input newline convention instead of silently
            # converting CRLF/LF while the content itself is sanitized.
            with io.TextIOWrapper(src_raw, encoding=codec, errors="strict", newline="") as src, io.TextIOWrapper(
                dst_raw, encoding=codec, errors="strict", newline=""
            ) as dst:
                for line_no, line in enumerate(src, 1):
                    sanitized, n, counts = sanitize_content(line, rules)
                    dst.write(sanitized)
                    total += n
                    for kind, count in counts.items():
                        audit_add(audit, audit_name, kind, count, f"Line {line_no}", sanitized)
        tmp_out.replace(fpath)
        return total
    except Exception:
        tmp_out.unlink(missing_ok=True)
        raise


def sanitize_xlsx_file(fpath: Path, rules: list[str], audit: dict, audit_name: str) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(fpath)
    total = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    sanitized, n, counts = sanitize_content(cell.value, rules)
                    cell.value = sanitized
                    total += n
                    for kind, count in counts.items():
                        audit_add(audit, audit_name, kind, count, f"{ws.title}!{cell.coordinate}", sanitized)
    wb.save(fpath)
    return total


def sanitize_xls_file(fpath: Path, rules: list[str], audit: dict, audit_name: str) -> int:
    import xlrd
    import xlwt

    rb = xlrd.open_workbook(fpath)
    wb = xlwt.Workbook()
    total = 0
    for sheet_idx in range(rb.nsheets):
        rs = rb.sheet_by_index(sheet_idx)
        ws = wb.add_sheet(rs.name[:31])
        for row_idx in range(rs.nrows):
            for col_idx in range(rs.ncols):
                value = rs.cell_value(row_idx, col_idx)
                if isinstance(value, str):
                    value, n, counts = sanitize_content(value, rules)
                    total += n
                    for kind, count in counts.items():
                        audit_add(
                            audit,
                            audit_name,
                            kind,
                            count,
                            f"{rs.name}!R{row_idx + 1}C{col_idx + 1}",
                            value,
                        )
                ws.write(row_idx, col_idx, value)
    wb.save(str(fpath))
    return total


def sanitize_file_if_supported(fpath: Path, rules: list[str], audit: dict, audit_name: str) -> int:
    """Sanitize supported documents, spreadsheets and text; fail closed otherwise."""
    ext = fpath.suffix.lower()
    if ext == ".pdf":
        return sanitize_pdf_document(
            fpath,
            rules,
            audit,
            audit_name,
            sanitize_content,
            classify_rule,
            audit_add,
        )
    if ext == ".docx":
        return sanitize_docx_document(
            fpath,
            rules,
            audit,
            audit_name,
            sanitize_content,
            audit_add,
        )
    if ext == ".xlsx":
        return sanitize_xlsx_file(fpath, rules, audit, audit_name)
    if ext == ".xls":
        return sanitize_xls_file(fpath, rules, audit, audit_name)

    if is_known_binary_extension(fpath):
        raise ValueError(
            f"Unsupported binary/document file type {ext or '(no extension)'}: {fpath.name}. "
            "A dedicated format sanitizer is required."
        )

    encoding_info = detect_text_encoding(fpath)
    if ext in cfg.TEXT_EXTS or ext in TEXT_LIKE_EXTS or encoding_info is not None:
        if encoding_info is None:
            raise ValueError(f"Unsupported text encoding or binary content: {fpath.name}")
        return sanitize_text_file(fpath, rules, audit, audit_name, encoding_info)

    raise ValueError(
        f"Unsupported or binary file type {ext or '(no extension)'}: {fpath.name}"
    )


def sanitize_filename_windows(name: str):
    bad_pattern = r'[<>:"/\\|?*\x00-\x1F]|[^\x20-\x7E]'
    sanitized = re.sub(bad_pattern, x_replacer, str(name)).rstrip(" .")
    if not sanitized or sanitized in (".", ".."):
        sanitized = "file"
    return sanitized, len(re.findall(bad_pattern, str(name)))


def get_archive_ext(name: str) -> str:
    lower = str(name).lower()
    for ext in cfg.COMPOUND_ARCHIVE_SUFFIXES:
        if lower.endswith(ext):
            return ext
    return Path(name).suffix.lower()


def is_archive_name(name: str) -> bool:
    return get_archive_ext(name) in cfg.ARCHIVE_EXTS


def archive_stem(name: str) -> str:
    text = str(name)
    ext = get_archive_ext(text)
    return text[:-len(ext)] if ext and text.lower().endswith(ext) else Path(text).stem


def sanitize_name_component(name: str, rules: list[str] | None = None, preserve_extension: bool = True):
    """Sanitize one path component while preserving a file/archive extension."""
    safe, format_count = sanitize_filename_windows(str(name))
    counts: dict[str, int] = {}
    if format_count:
        counts["Filename Format"] = format_count
    if not rules:
        return safe, format_count, counts

    ext = ""
    body = safe
    if preserve_extension:
        ext = get_archive_ext(safe) if is_archive_name(safe) else Path(safe).suffix
        if ext and safe.lower().endswith(ext.lower()):
            body = safe[: -len(ext)]
    body, rule_count, rule_counts = sanitize_content(body, rules)
    for kind, count in rule_counts.items():
        counts[kind] = counts.get(kind, 0) + count
    safe = (body + ext).rstrip(" .") or ("file" + ext)
    return safe, format_count + rule_count, counts


def sanitize_relative_path(relative_path: str, rules: list[str] | None = None):
    """Return a safe relative path, replacement count and match categories."""
    raw = str(relative_path or "").replace("\\", "/")
    raw_parts = [p for p in PurePosixPath(raw).parts if p not in ("", ".", "..", "/")]
    if not raw_parts:
        raw_parts = ["file"]
    parts: list[str] = []
    replacements = 0
    counts: dict[str, int] = {}
    for idx, part in enumerate(raw_parts):
        safe, n, component_counts = sanitize_name_component(
            part, rules, preserve_extension=(idx == len(raw_parts) - 1)
        )
        replacements += n
        for kind, count in component_counts.items():
            counts[kind] = counts.get(kind, 0) + count
        parts.append(safe)
    return Path(*parts), replacements, counts


def is_junk_mac_file(path: Path) -> bool:
    text = str(path).replace("\\", "/")
    return "/__MACOSX" in text or text.endswith("/.DS_Store") or path.name == ".DS_Store"


def _safe_target(root: Path, member_name: str) -> Path:
    clean = str(member_name).replace("\\", "/")
    pure = PurePosixPath(clean)
    if pure.is_absolute() or ".." in pure.parts:
        raise ValueError(f"Unsafe archive member path: {member_name}")
    target = (root / Path(*[p for p in pure.parts if p not in ("", ".")])).resolve()
    root_resolved = root.resolve()
    if target != root_resolved and root_resolved not in target.parents:
        raise ValueError(f"Archive member escapes extraction directory: {member_name}")
    return target


def _validate_extracted_tree(root: Path) -> None:
    count = 0
    total = 0
    for path in root.rglob("*"):
        if path.is_symlink():
            raise ValueError("Symlinks are not permitted in extracted archives")
        if path.is_file():
            count += 1
            total += path.stat().st_size
            if count > cfg.MAX_EXTRACTED_FILES:
                raise ValueError(f"Archive exceeds {cfg.MAX_EXTRACTED_FILES} extracted files")
            if total > cfg.MAX_EXTRACTED_MB * 1024 * 1024:
                raise ValueError(f"Archive exceeds {cfg.MAX_EXTRACTED_MB} MB extracted size")


def _extract_zip(path: Path, dest: Path) -> None:
    with zipfile.ZipFile(path) as zf:
        members = zf.infolist()
        if len(members) > cfg.MAX_EXTRACTED_FILES:
            raise ValueError("ZIP contains too many files")
        total = sum(m.file_size for m in members)
        if total > cfg.MAX_EXTRACTED_MB * 1024 * 1024:
            raise ValueError("ZIP extracted size exceeds configured limit")
        for member in members:
            mode = (member.external_attr >> 16) & 0o170000
            if mode == stat.S_IFLNK:
                raise ValueError("Symlinks are not permitted in ZIP archives")
            target = _safe_target(dest, member.filename)
            if member.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(member) as src, open(target, "wb") as dst:
                shutil.copyfileobj(src, dst)


def _extract_tar(path: Path, dest: Path) -> None:
    with tarfile.open(path, "r:*") as tf:
        members = tf.getmembers()
        files = [m for m in members if m.isfile()]
        if len(files) > cfg.MAX_EXTRACTED_FILES:
            raise ValueError("TAR contains too many files")
        total = sum(m.size for m in files)
        if total > cfg.MAX_EXTRACTED_MB * 1024 * 1024:
            raise ValueError("TAR extracted size exceeds configured limit")
        for member in members:
            if member.issym() or member.islnk() or member.isdev():
                raise ValueError("Links and device files are not permitted in archives")
            target = _safe_target(dest, member.name)
            if member.isdir():
                target.mkdir(parents=True, exist_ok=True)
            elif member.isfile():
                target.parent.mkdir(parents=True, exist_ok=True)
                src = tf.extractfile(member)
                if src is not None:
                    with src, open(target, "wb") as dst:
                        shutil.copyfileobj(src, dst)


def _extract_single_stream(path: Path, dest: Path, kind: str) -> None:
    output_name = path.stem or "content"
    target = _safe_target(dest, output_name)
    opener = {".gz": gzip.open, ".bz2": bz2.open, ".xz": lzma.open}[kind]
    target.parent.mkdir(parents=True, exist_ok=True)
    total = 0
    with opener(path, "rb") as src, open(target, "wb") as dst:
        while True:
            chunk = src.read(1024 * 1024)
            if not chunk:
                break
            total += len(chunk)
            if total > cfg.MAX_EXTRACTED_MB * 1024 * 1024:
                raise ValueError("Compressed stream exceeds extracted size limit")
            dst.write(chunk)


def extract_archive(upload_path: Path, dest: Path, depth: int = 0) -> None:
    if depth > cfg.MAX_ARCHIVE_DEPTH:
        raise ValueError(f"Archive nesting exceeds maximum depth {cfg.MAX_ARCHIVE_DEPTH}")
    upload_path = Path(upload_path)
    dest = Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    ext = get_archive_ext(upload_path.name)

    if ext == ".zip":
        _extract_zip(upload_path, dest)
    elif ext in (".tar", ".tar.gz", ".tgz", ".tar.bz2", ".tbz2", ".tar.xz"):
        _extract_tar(upload_path, dest)
    elif ext in (".gz", ".bz2", ".xz"):
        _extract_single_stream(upload_path, dest, ext)
    else:
        exe = cfg.ensure_7zip(False)
        if not exe:
            raise RuntimeError(
                f"{ext or 'This archive'} requires 7-Zip, but no local/bundled 7-Zip executable was found"
            )
        result = subprocess.run(
            [str(exe), "x", str(upload_path), f"-o{dest}", "-y"],
            capture_output=True,
            text=True,
            timeout=cfg.ARCHIVE_TIMEOUT_SECONDS,
        )
        if result.returncode != 0:
            raise RuntimeError((result.stderr or result.stdout or "7-Zip extraction failed")[-1000:])
    _validate_extracted_tree(dest)


def repack_directory(root: Path, out_base: Path, original_ext: str) -> Path:
    root = Path(root)
    original_ext = original_ext.lower()
    if original_ext == ".tar":
        return Path(shutil.make_archive(str(out_base), "tar", root_dir=root))
    if original_ext in (".tar.gz", ".tgz"):
        return Path(shutil.make_archive(str(out_base), "gztar", root_dir=root))
    if original_ext in (".tar.bz2", ".tbz2"):
        return Path(shutil.make_archive(str(out_base), "bztar", root_dir=root))
    if original_ext == ".tar.xz":
        return Path(shutil.make_archive(str(out_base), "xztar", root_dir=root))

    files = [p for p in root.rglob("*") if p.is_file()]
    if original_ext in (".gz", ".bz2", ".xz") and len(files) == 1:
        src = files[0]
        out = Path(str(out_base) + original_ext)
        opener = {".gz": gzip.open, ".bz2": bz2.open, ".xz": lzma.open}[original_ext]
        with open(src, "rb") as inp, opener(out, "wb") as dst:
            shutil.copyfileobj(inp, dst)
        return out

    if original_ext == ".7z":
        exe = cfg.ensure_7zip(False)
        if exe:
            out = Path(str(out_base) + ".7z")
            result = subprocess.run(
                [str(exe), "a", "-t7z", str(out), "."],
                cwd=str(root),
                capture_output=True,
                text=True,
                timeout=cfg.ARCHIVE_TIMEOUT_SECONDS,
            )
            if result.returncode == 0 and out.exists():
                return out

    return Path(shutil.make_archive(str(out_base), "zip", root_dir=root))


def _unique_sanitized_name(path: Path, used: set[str], rules: list[str] | None = None, is_dir: bool = False):
    safe, n, counts = sanitize_name_component(path.name, rules, preserve_extension=not is_dir)
    base = safe
    counter = 1
    while safe in used or ((path.parent / safe).exists() and (path.parent / safe) != path):
        stem, ext = os.path.splitext(base)
        safe = f"{stem}_{counter}{ext}"
        counter += 1
    used.add(safe)
    target = path.parent / safe
    if target != path:
        path.rename(target)
    return target, n, counts


def sanitize_directory(root: Path, rules: list[str], audit: dict, audit_prefix: str, depth: int = 0) -> int:
    if depth > cfg.MAX_ARCHIVE_DEPTH:
        raise ValueError(f"Archive nesting exceeds maximum depth {cfg.MAX_ARCHIVE_DEPTH}")
    root = Path(root)
    total = 0

    for dirpath, _, filenames in os.walk(root, topdown=False):
        directory = Path(dirpath)
        used: set[str] = set()
        for name in filenames:
            path = directory / name
            if not path.exists() or is_junk_mac_file(path):
                if path.exists():
                    path.unlink(missing_ok=True)
                continue
            relative_before = path.relative_to(root).as_posix()
            audit_name = f"{audit_prefix}/{relative_before}" if audit_prefix else relative_before

            if is_archive_name(path.name):
                nested_dir = Path(tempfile.mkdtemp(prefix="nested_", dir=cfg.TEMP))
                try:
                    ext = get_archive_ext(path.name)
                    extract_archive(path, nested_dir, depth + 1)
                    total += sanitize_directory(nested_dir, rules, audit, audit_name, depth + 1)
                    out_base = nested_dir.parent / f"repacked_{uuid.uuid4().hex[:8]}"
                    repacked = repack_directory(nested_dir, out_base, ext)
                    actual_ext = get_archive_ext(repacked.name)
                    replacement = path
                    if actual_ext and actual_ext != ext:
                        replacement = path.with_name(archive_stem(path.name) + actual_ext)
                    path.unlink(missing_ok=True)
                    shutil.move(str(repacked), str(replacement))
                    path = replacement
                finally:
                    shutil.rmtree(nested_dir, ignore_errors=True)
            else:
                total += sanitize_file_if_supported(path, rules, audit, audit_name)

            path, fname_repl, name_counts = _unique_sanitized_name(path, used, rules, is_dir=False)
            if fname_repl:
                total += fname_repl
                safe_rel = path.relative_to(root).as_posix()
                for kind, count in name_counts.items():
                    audit_add(audit, audit_name, f"Filename: {kind}", count, "Filename", safe_rel)

    for dirpath, dirnames, _ in os.walk(root, topdown=False):
        directory = Path(dirpath)
        used = {p.name for p in directory.iterdir() if p.is_file()}
        for name in dirnames:
            path = directory / name
            if not path.exists() or is_junk_mac_file(path):
                if path.exists():
                    shutil.rmtree(path, ignore_errors=True)
                continue
            old_rel = path.relative_to(root).as_posix()
            new_path, n, name_counts = _unique_sanitized_name(path, used, rules, is_dir=True)
            if n:
                total += n
                label = f"{audit_prefix}/{old_rel}" if audit_prefix else old_rel
                for kind, count in name_counts.items():
                    audit_add(audit, label, f"Directory: {kind}", count, "Directory", new_path.name)
    return total
