#!/usr/bin/env python3
"""Authenticated staging acceptance test for SOC Data Sanitization Platform v2.1."""
from __future__ import annotations

import argparse
import getpass
import io
import json
import re
import ssl
import sys
import time
import uuid
import zipfile
from http.cookiejar import CookieJar
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import HTTPCookieProcessor, HTTPSHandler, Request, build_opener

import openpyxl


def die(message: str) -> None:
    print(f"FAIL: {message}", file=sys.stderr)
    raise SystemExit(1)


def main() -> int:
    parser = argparse.ArgumentParser(description="Run authenticated staging acceptance checks")
    parser.add_argument("--url", default="https://127.0.0.1:8443", help="Application base URL")
    parser.add_argument("--username", required=True, help="Existing staging username")
    parser.add_argument("--role", required=True, choices=["admin", "analyst"], help="Expected role")
    parser.add_argument("--ca", help="CA certificate file. Omit for staging self-signed TLS.")
    parser.add_argument("--timeout", type=int, default=30, help="Job completion timeout in seconds")
    args = parser.parse_args()

    base = args.url.rstrip("/")
    password = getpass.getpass(f"Password for {args.username}: ")
    if not password:
        die("password cannot be empty")

    if args.ca:
        context = ssl.create_default_context(cafile=args.ca)
    else:
        context = ssl._create_unverified_context()  # staging only; production should pass --ca

    jar = CookieJar()
    opener = build_opener(HTTPCookieProcessor(jar), HTTPSHandler(context=context))

    def request(path: str, *, method: str = "GET", data: bytes | None = None, headers: dict | None = None, allow: tuple[int, ...] = (200,)):
        req = Request(base + path, data=data, method=method, headers=headers or {})
        try:
            response = opener.open(req, timeout=15)
            status = response.status
            body = response.read()
            hdrs = response.headers
        except HTTPError as exc:
            status = exc.code
            body = exc.read()
            hdrs = exc.headers
        except URLError as exc:
            die(f"connection failed: {exc}")
        if status not in allow:
            preview = body.decode("utf-8", errors="replace")[:500]
            die(f"{method} {path} returned HTTP {status}: {preview}")
        return status, hdrs, body

    print("SOC Data Sanitization Platform v2.1 - Staging Acceptance")
    print("=" * 64)
    print(f"Target:   {base}")
    print(f"User:     {args.username}")
    print(f"Role:     {args.role}")

    # Unauthenticated health endpoint.
    _, _, body = request("/healthz")
    healthz = json.loads(body)
    if healthz.get("status") != "ok":
        die("/healthz did not report ok")
    print("PASS: unauthenticated service health")

    # Login page and CSRF token.
    _, login_headers, login_body = request("/login")
    login_html = login_body.decode("utf-8", errors="replace")
    match = re.search(r'name="csrf_token" value="([^"]+)"', login_html)
    if not match:
        die("login CSRF token not found")
    csrf_login = match.group(1)

    form = urlencode(
        {
            "username": args.username,
            "password": password,
            "csrf_token": csrf_login,
            "next": "/",
        }
    ).encode()
    # urllib follows the successful login redirect automatically.
    _, root_headers, root_body = request(
        "/login",
        method="POST",
        data=form,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
    )
    root_html = root_body.decode("utf-8", errors="replace")
    if "SOC Data Sanitization" not in root_html and "Sanitization Platform" not in root_html:
        die("login did not reach the application dashboard")
    print("PASS: authenticated login and dashboard")

    # Required browser security headers.
    required_headers = {
        "X-Frame-Options": "DENY",
        "X-Content-Type-Options": "nosniff",
        "Referrer-Policy": "no-referrer",
    }
    for name, expected in required_headers.items():
        actual = root_headers.get(name)
        if actual != expected:
            die(f"security header {name} expected {expected!r}, found {actual!r}")
    csp = root_headers.get("Content-Security-Policy", "")
    if "default-src 'self'" not in csp or "frame-ancestors 'none'" not in csp:
        die("Content-Security-Policy is missing expected restrictions")
    print("PASS: browser security headers")

    _, _, session_body = request("/api/session")
    session_data = json.loads(session_body)
    if session_data.get("username") != args.username.lower():
        die(f"session username mismatch: {session_data.get('username')!r}")
    if session_data.get("role") != args.role:
        die(f"expected role {args.role}, got {session_data.get('role')!r}")
    csrf = session_data.get("csrf_token")
    if not csrf:
        die("authenticated CSRF token missing")
    print("PASS: authenticated session and expected RBAC role")

    _, _, auth_health_body = request("/health")
    auth_health = json.loads(auth_health_body)
    if auth_health.get("status") != "ok" or auth_health.get("version") != "2.1":
        die(f"authenticated /health unexpected: {auth_health}")
    print("PASS: authenticated application health")

    # RBAC: admins can read rules; analysts must be forbidden.
    if args.role == "admin":
        _, _, rules_body = request("/bad_words")
        rules_data = json.loads(rules_body)
        if not isinstance(rules_data.get("lines"), list):
            die("admin rules endpoint returned unexpected payload")
        print("PASS: administrator Rules Library access")
    else:
        status, _, _ = request("/bad_words", allow=(403,))
        if status != 403:
            die("analyst unexpectedly gained Rules Library access")
        print("PASS: analyst Rules Library restriction")

    # Non-destructive pipeline test. It creates a job but does not alter rules.
    boundary = "----SanitAcceptance" + uuid.uuid4().hex
    filename = "staging_acceptance.txt"
    content = b"SOC SANITIZATION STAGING ACCEPTANCE TEST\n"
    multipart = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        "Content-Type: text/plain\r\n\r\n"
    ).encode() + content + (
        f"\r\n--{boundary}\r\n"
        'Content-Disposition: form-data; name="relative_path"\r\n\r\n'
        f"Acceptance/{filename}\r\n"
        f"--{boundary}--\r\n"
    ).encode()
    _, _, upload_body = request(
        "/upload",
        method="POST",
        data=multipart,
        headers={
            "Content-Type": f"multipart/form-data; boundary={boundary}",
            "X-CSRF-Token": csrf,
        },
    )
    upload = json.loads(upload_body)
    job_id = upload.get("job_id")
    if not job_id:
        die(f"upload did not return a job id: {upload}")
    print(f"PASS: upload accepted as job {job_id}")

    deadline = time.time() + args.timeout
    result = None
    while time.time() < deadline:
        _, _, status_body = request(f"/status/{job_id}")
        result = json.loads(status_body)
        if result.get("status") in {"Done", "Failed"}:
            break
        time.sleep(0.25)
    if not result or result.get("status") != "Done":
        die(f"pipeline job did not complete successfully: {result}")
    if not result.get("download") or not result.get("report_download"):
        die("completed job is missing package/report download links")
    print("PASS: sanitization processing reached Done")

    _, _, package_bytes = request(result["download"])
    try:
        with zipfile.ZipFile(io.BytesIO(package_bytes)) as package:
            names = package.namelist()
            expected = f"Sanitized_Files/Acceptance/{filename}"
            if expected not in names:
                die(f"sanitized package missing {expected}; entries={names}")
            if package.read(expected) != content:
                die("benign staging file changed unexpectedly")
    except zipfile.BadZipFile:
        die("sanitized package is not a valid ZIP file")
    print("PASS: sanitized output package download")

    _, _, report_bytes = request(result["report_download"])
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(report_bytes), data_only=True)
    except Exception as exc:
        die(f"Excel audit report could not be opened: {exc}")
    if not {"Summary", "Audit"}.issubset(set(workbook.sheetnames)):
        die(f"Excel report missing expected sheets: {workbook.sheetnames}")
    print("PASS: Excel audit report download and structure")

    _, _, jobs_body = request("/jobs?limit=20")
    jobs = json.loads(jobs_body).get("jobs", [])
    found = next((item for item in jobs if item.get("job_id") == job_id), None)
    if not found or found.get("status") != "Done":
        die("completed staging job not found in persistent job history")
    if found.get("created_by") != args.username.lower():
        die(f"job attribution mismatch: {found.get('created_by')!r}")
    print("PASS: persistent job history and user attribution")

    # Logout must require a valid CSRF token.
    bad_form = urlencode({"csrf_token": "invalid"}).encode()
    status, _, _ = request(
        "/logout",
        method="POST",
        data=bad_form,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        allow=(403,),
    )
    if status != 403:
        die("logout accepted an invalid CSRF token")
    good_form = urlencode({"csrf_token": csrf}).encode()
    _, _, logout_body = request(
        "/logout",
        method="POST",
        data=good_form,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
    )
    if b"Sign In" not in logout_body and b"sign in" not in logout_body.lower():
        die("logout did not return to the sign-in page")
    print("PASS: CSRF-protected logout")

    print("-" * 64)
    print("RESULT: PASSED")
    print(f"Staging user {args.username!r} validated successfully as {args.role}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
