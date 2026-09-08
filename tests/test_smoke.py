import io
import re
import time
import zipfile

import openpyxl
import pytest

from sanitization_v2 import auth
from sanitization_v2 import config as cfg
from sanitization_v2 import rules as rules_module
from sanitization_v2.app import app

ADMIN_PASSWORD = "Admin-Testing-Password-123!"
ANALYST_PASSWORD = "Analyst-Testing-Password-123!"


@pytest.fixture(autouse=True)
def isolated_runtime(tmp_path, monkeypatch):
    paths = {
        "UPLOAD": tmp_path / "uploads",
        "OUTPUT": tmp_path / "outputs",
        "LOGS": tmp_path / "logs",
        "TEMP": tmp_path / "temp",
        "DATA": tmp_path / "data",
    }
    for name, path in paths.items():
        path.mkdir(parents=True, exist_ok=True)
        monkeypatch.setattr(cfg, name, path)

    monkeypatch.setattr(cfg, "DB_PATH", paths["DATA"] / "jobs.db")
    rules_file = tmp_path / "bad_words.txt"
    rules_file.write_text("TESTSECRET\n", encoding="utf-8")
    monkeypatch.setattr(cfg, "RULES_FILE", str(rules_file))
    monkeypatch.setattr(cfg, "SEVEN_ZIP_EXE", None)

    cfg.SESSIONS.clear()
    cfg.init_db()
    auth.create_user("admin", ADMIN_PASSWORD, role="admin")
    auth.create_user("analyst", ANALYST_PASSWORD, role="analyst")
    rules_module.reload_bad_patterns()
    app.config.update(TESTING=True, SESSION_COOKIE_SECURE=False)
    yield
    cfg.SESSIONS.clear()


def login(client, username="admin", password=ADMIN_PASSWORD):
    page = client.get("/login")
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match, html
    response = client.post(
        "/login",
        data={"username": username, "password": password, "csrf_token": match.group(1), "next": "/"},
        follow_redirects=False,
    )
    assert response.status_code == 302
    session_response = client.get("/api/session")
    assert session_response.status_code == 200
    return session_response.get_json()["csrf_token"]


def wait_for_job(client, job_id, timeout=15):
    deadline = time.time() + timeout
    last = None
    while time.time() < deadline:
        response = client.get(f"/status/{job_id}")
        assert response.status_code == 200
        last = response.get_json()
        if last["status"] in ("Done", "Failed"):
            return last
        time.sleep(0.05)
    pytest.fail(f"job {job_id} did not finish; last status={last}")


def test_authentication_ui_roles_health_and_security_headers():
    client = app.test_client()

    root = client.get("/")
    assert root.status_code == 302
    assert "/login" in root.headers["Location"]
    assert client.get("/jobs").status_code == 401
    assert client.get("/healthz").get_json()["status"] == "ok"

    csrf = login(client)
    assert csrf
    response = client.get("/")
    assert response.status_code == 200
    html = response.get_data(as_text=True)
    for required in (
        'data-view="dashboard"',
        'data-view="sanitize"',
        'data-view="history"',
        'data-view="rules"',
        'id="browse"',
        'id="browseFolder"',
        'id="clear"',
        'id="start"',
        'id="refreshJobs"',
        'id="exportJobs"',
        'id="addRule"',
        'id="saveRules"',
        'id="clearRules"',
        'id="replaceRules"',
        'id="healthBtn"',
    ):
        assert required in html
    assert "Administrator" in html
    assert response.headers["X-Frame-Options"] == "DENY"
    assert "default-src 'self'" in response.headers["Content-Security-Policy"]

    health = client.get("/health")
    assert health.status_code == 200
    data = health.get_json()
    assert data["status"] == "ok"
    assert data["version"] == "2.1"
    assert data["offline"] is True
    assert data["role"] == "admin"

    analyst = app.test_client()
    analyst_csrf = login(analyst, "analyst", ANALYST_PASSWORD)
    analyst_html = analyst.get("/").get_data(as_text=True)
    assert 'data-view="rules"' not in analyst_html
    assert 'id="addRule"' not in analyst_html
    assert analyst.get("/bad_words").status_code == 403
    blocked = analyst.post(
        "/rules/save",
        json={"lines": ["TESTSECRET"]},
        headers={"X-CSRF-Token": analyst_csrf},
    )
    assert blocked.status_code == 403


def test_login_failure_and_logout_csrf():
    client = app.test_client()
    page = client.get("/login").get_data(as_text=True)
    token = re.search(r'name="csrf_token" value="([^"]+)"', page).group(1)
    response = client.post(
        "/login",
        data={"username": "admin", "password": "wrong-password", "csrf_token": token},
    )
    assert response.status_code == 401
    assert "Invalid username or password" in response.get_data(as_text=True)

    csrf = login(client)
    assert client.post("/logout", data={"csrf_token": "bad"}).status_code == 403
    response = client.post("/logout", data={"csrf_token": csrf})
    assert response.status_code == 302
    assert client.get("/jobs").status_code == 401


def test_rule_management_endpoints_require_admin_and_csrf():
    client = app.test_client()
    csrf = login(client)

    assert client.post("/rules/save", json={"lines": ["TESTSECRET"]}).status_code == 403

    response = client.post(
        "/rules/save",
        json={"lines": ["TESTSECRET"]},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200

    response = client.post(
        "/rules/add",
        json={"rule": "NEWSECRET"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200
    assert response.get_json()["count"] == 2

    response = client.post(
        "/rules/edit",
        json={"index": 1, "line": "CHANGEDSECRET"},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200

    lines = client.get("/bad_words").get_json()["lines"]
    assert lines == ["TESTSECRET", "CHANGEDSECRET"]

    response = client.post(
        "/rules/remove",
        json={"index": 1},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200
    assert client.get("/bad_words").get_json()["lines"] == ["TESTSECRET"]

    response = client.post(
        "/rules/add",
        json={"rule": "["},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 400
    assert "invalid regex" in response.get_json()["error"]


def test_multifile_folder_package_report_and_history_as_analyst():
    client = app.test_client()
    csrf = login(client, "analyst", ANALYST_PASSWORD)
    payload = {
        "file": [
            (io.BytesIO(b"alpha TESTSECRET omega\n"), "one.txt"),
            (io.BytesIO(b"second TESTSECRET value\n"), "two.log"),
        ],
        "relative_path": ["Case-A/one.txt", "Case-A/sub/two.log"],
    }
    response = client.post(
        "/upload",
        data=payload,
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200
    job_id = response.get_json()["job_id"]

    result = wait_for_job(client, job_id)
    assert result["status"] == "Done", result
    assert result["total_replacements"] >= 2
    assert result["download"]
    assert result["report_download"]

    package_response = client.get(result["download"])
    assert package_response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(package_response.data)) as package:
        names = package.namelist()
        assert "Sanitized_Files/Case-A/one.txt" in names
        assert "Sanitized_Files/Case-A/sub/two.log" in names
        assert any(name.endswith(f"Sanitization_Report_{job_id}.xlsx") for name in names)
        one = package.read("Sanitized_Files/Case-A/one.txt").decode("utf-8")
        two = package.read("Sanitized_Files/Case-A/sub/two.log").decode("utf-8")
        assert "TESTSECRET" not in one
        assert "TESTSECRET" not in two
        assert "X" in one and "X" in two

    report_response = client.get(result["report_download"])
    assert report_response.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(report_response.data), data_only=True)
    assert {"Summary", "Audit"}.issubset(wb.sheetnames)
    report_text = "\n".join(
        str(value or "")
        for sheet in wb.worksheets
        for row in sheet.iter_rows(values_only=True)
        for value in row
    )
    assert "TESTSECRET" not in report_text
    assert "Total Replacements" in report_text

    jobs = client.get("/jobs?limit=20").get_json()["jobs"]
    found = next(item for item in jobs if item["job_id"] == job_id)
    assert found["status"] == "Done"
    assert found["created_by"] == "analyst"
    assert found["download_available"] is True
    assert found["report_available"] is True


def test_archive_path_traversal_becomes_failed_job():
    client = app.test_client()
    csrf = login(client, "analyst", ANALYST_PASSWORD)
    archive_bytes = io.BytesIO()
    with zipfile.ZipFile(archive_bytes, "w") as zf:
        zf.writestr("../escape.txt", "TESTSECRET")
    archive_bytes.seek(0)

    response = client.post(
        "/upload",
        data={
            "file": [(archive_bytes, "unsafe.zip")],
            "relative_path": ["unsafe.zip"],
        },
        content_type="multipart/form-data",
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200
    job_id = response.get_json()["job_id"]
    result = wait_for_job(client, job_id)
    assert result["status"] == "Failed"
    assert "Unsafe archive member path" in result["error"]

    jobs = client.get("/jobs?limit=20").get_json()["jobs"]
    found = next(item for item in jobs if item["job_id"] == job_id)
    assert found["status"] == "Failed"
    assert found["created_by"] == "analyst"
    assert found["error"]
