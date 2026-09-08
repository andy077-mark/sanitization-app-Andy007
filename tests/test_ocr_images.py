import io
import zipfile
from pathlib import Path

import pytest

from sanitization_v2.ocr_sanitizer import find_tesseract, ocr_image_matches
from sanitization_v2.sanitize import classify_rule, sanitize_file_if_supported

RULES = ["TESTSECRET"]


def _font():
    from PIL import ImageFont

    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, 58)
    pytest.skip("DejaVu font not available for deterministic OCR test")


def _make_test_image(path: Path, add_metadata: bool = False):
    from PIL import Image, ImageDraw, PngImagePlugin

    image = Image.new("RGB", (1100, 260), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 80), "TESTSECRET", font=_font(), fill="black")
    if add_metadata:
        meta = PngImagePlugin.PngInfo()
        meta.add_text("Comment", "TESTSECRET")
        image.save(path, format="PNG", pnginfo=meta)
    else:
        image.save(path, format="PNG")


def _sanitize(path: Path):
    audit = {}
    count = sanitize_file_if_supported(path, RULES, audit, path.name)
    return count, audit


def _ocr_has_secret(image) -> bool:
    _, matches = ocr_image_matches(image, RULES, classify_rule)
    return bool(matches)


@pytest.mark.skipif(find_tesseract() is None, reason="Tesseract OCR not installed")
def test_png_ocr_redacts_sensitive_text_and_strips_metadata(tmp_path):
    from PIL import Image

    path = tmp_path / "screenshot.png"
    _make_test_image(path, add_metadata=True)

    count, audit = _sanitize(path)

    assert count >= 1
    with Image.open(path) as sanitized:
        sanitized.load()
        assert not _ocr_has_secret(sanitized.copy())
        assert "Comment" not in sanitized.info
    assert any("Image OCR" in item["location"] for item in audit.values())


@pytest.mark.skipif(find_tesseract() is None, reason="Tesseract OCR not installed")
def test_scanned_pdf_uses_ocr_and_permanent_image_redaction(tmp_path):
    import pymupdf as fitz
    from PIL import Image

    image_path = tmp_path / "scan.png"
    _make_test_image(image_path)

    pdf_path = tmp_path / "scan.pdf"
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(fitz.Rect(40, 220, 572, 346), filename=str(image_path))
    doc.save(pdf_path)
    doc.close()

    count, audit = _sanitize(pdf_path)

    assert count >= 1
    with fitz.open(pdf_path) as sanitized:
        page = sanitized[0]
        pix = page.get_pixmap(dpi=200, alpha=False)
        raster = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        assert not _ocr_has_secret(raster)
    assert any("Page 1 image" in item["location"] and "OCR" in item["location"] for item in audit.values())


@pytest.mark.skipif(find_tesseract() is None, reason="Tesseract OCR not installed")
def test_docx_embedded_image_is_ocr_sanitized(tmp_path):
    from docx import Document
    from PIL import Image

    image_path = tmp_path / "embedded.png"
    _make_test_image(image_path)

    docx_path = tmp_path / "evidence.docx"
    doc = Document()
    doc.add_picture(str(image_path))
    doc.save(docx_path)

    count, audit = _sanitize(docx_path)

    assert count >= 1
    Document(docx_path)  # package remains valid
    with zipfile.ZipFile(docx_path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert media
        data = archive.read(media[0])
    with Image.open(io.BytesIO(data)) as sanitized_image:
        sanitized_image.load()
        assert not _ocr_has_secret(sanitized_image.copy())
    assert any("DOCX image" in item["location"] and "OCR" in item["location"] for item in audit.values())


def test_image_fails_closed_when_ocr_engine_is_unavailable(tmp_path, monkeypatch):
    path = tmp_path / "screenshot.png"
    _make_test_image(path)
    monkeypatch.setenv("SANIT_TESSERACT_BINARY", str(tmp_path / "missing-tesseract"))

    with pytest.raises(ValueError, match="OCR engine unavailable"):
        _sanitize(path)
