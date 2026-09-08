import zipfile

import pytest

from sanitization_v2.sanitize import sanitize_file_if_supported

RULES = ["TESTSECRET", r"\b10\.20\.30\.40\b"]


def sanitize(path):
    audit = {}
    count = sanitize_file_if_supported(path, RULES, audit, path.name)
    return count, audit


def test_pdf_text_is_permanently_redacted_and_reopens(tmp_path):
    import pymupdf as fitz

    path = tmp_path / "evidence.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "User=TESTSECRET IP=10.20.30.40")
    doc.set_metadata({"title": "Case TESTSECRET"})
    doc.save(path)
    doc.close()

    count, audit = sanitize(path)

    assert count >= 3
    with fitz.open(path) as result:
        text = "\n".join(page.get_text("text") for page in result)
        metadata = result.metadata or {}
    assert "TESTSECRET" not in text
    assert "10.20.30.40" not in text
    assert "TESTSECRET" not in (metadata.get("title") or "")
    assert "X" in text
    assert any(item["location"].startswith("Page 1") for item in audit.values())


def test_pdf_image_only_page_fails_closed_without_ocr(tmp_path):
    import pymupdf as fitz

    path = tmp_path / "scanned.pdf"
    doc = fitz.open()
    page = doc.new_page(width=100, height=100)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10), False)
    pix.clear_with(255)
    page.insert_image(page.rect, pixmap=pix)
    doc.save(path)
    doc.close()

    with pytest.raises(ValueError, match="image-only|OCR"):
        sanitize(path)


def test_docx_sanitizes_text_across_runs_tables_header_and_metadata(tmp_path):
    from docx import Document

    path = tmp_path / "incident.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("User=TEST")
    paragraph.add_run("SECRET")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "IP=10.20.30.40"
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Header TESTSECRET"
    doc.core_properties.title = "Case TESTSECRET"
    doc.save(path)

    count, audit = sanitize(path)
    assert count >= 4

    result = Document(path)
    body_text = "\n".join(p.text for p in result.paragraphs)
    table_text = "\n".join(
        cell.text for table in result.tables for row in table.rows for cell in row.cells
    )
    header_text = "\n".join(
        p.text for section in result.sections for p in section.header.paragraphs
    )
    assert "TESTSECRET" not in body_text
    assert "10.20.30.40" not in table_text
    assert "TESTSECRET" not in header_text
    assert "TESTSECRET" not in (result.core_properties.title or "")
    assert "XXXXXXXXXX" in body_text
    assert any("Document paragraph" in item["location"] for item in audit.values())


def test_docx_embedded_content_fails_closed(tmp_path):
    from docx import Document

    path = tmp_path / "embedded.docx"
    doc = Document()
    doc.add_paragraph("User=TESTSECRET")
    doc.save(path)

    rebuilt = tmp_path / "embedded_rebuilt.docx"
    with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(rebuilt, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info.filename))
        dst.writestr("word/embeddings/object1.bin", b"TESTSECRET")
    rebuilt.replace(path)

    with pytest.raises(ValueError, match="embedded/OLE"):
        sanitize(path)
